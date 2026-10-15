"""
Update Ziyo_Product_UX_Gap_Analysis.pptx with research-backed content.
Sources: Aksana (UX Designer) interview, Sirojiddin (Events Head) interview,
         Ziyo_App_Usage_Analysis.docx, prototype screenshot review.
"""
from pptx import Presentation

PPTX_PATH = '/Users/I039603/Downloads/ZiyoForum/Ziyo_Product_UX_Gap_Analysis.pptx'


def set_text(shape, *lines):
    """Replace text in a shape, one string per existing paragraph. Preserves run formatting."""
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    paras = tf.paragraphs
    for i, line in enumerate(lines):
        if i >= len(paras):
            break
        para = paras[i]
        if para.runs:
            para.runs[0].text = line
            for run in para.runs[1:]:
                run.text = ''


def show(label, shape):
    print(f"  {label}: \"{shape.text_frame.text.strip()[:90]}\"")


prs = Presentation(PPTX_PATH)

# ─────────────────────────────────────────────
# SLIDE 1  —  Cover / Stats
# ─────────────────────────────────────────────
s1 = prs.slides[0]
sh1 = s1.shapes

# Screens mapped: 18+ → 26+
set_text(sh1[11], '26+')

# Feature gaps: 6 → 9  (added: Skills & CV, Mock Tests, Jobs & Internships)
set_text(sh1[15], '9')

# Key finding — anchor to real uninstall metric
set_text(sh1[24],
    'The prototype demonstrates a substantially richer product — 26+ screens with gamification, skills & CV, jobs, '
    'and social features. The current app records an 83.8 % Android uninstall rate, 5 unanswered 1-star reviews, '
    'and a UI unchanged for over 12 months. Closing this gap requires an immediate fix sprint followed by a '
    'structured three-phase feature roadmap.'
)

print('Slide 1 updated.')

# ─────────────────────────────────────────────
# SLIDE 2  —  Executive Summary
# ─────────────────────────────────────────────
s2 = prs.slides[1]
sh2 = s2.shapes

# Current app — Basic Content Hub description
set_text(sh2[8],
    'Events, courses, grants, and job listings in a grey, minimal list UI. '
    'Interface unchanged for 12+ months — designed by outsourced team; '
    'UX Designer cannot iterate without file access. '
    '"The app is pretty absent — there\'s nothing in it." — Aksana, UX Designer'
)

# Current app — Limited Personalisation description
set_text(sh2[11],
    'No onboarding or interest-selection flow; same content feed for all 297 active Android users regardless '
    'of role, city, or interest area.'
)

# Current app — No Gamification description
set_text(sh2[14],
    'ZiyoCoins concept exists but is invisible: no balance shown, no earning events, no shop. '
    '83.8 % Android uninstall rate reflects zero daily-return incentive.'
)

# Current app — Single Content Format description
set_text(sh2[17],
    'Text/image cards only; no video, no audio. '
    'Android image distortion bug ("pictures look smashed") reported to developer for months — still unfixed.'
)

# Current app — Thin Social Features description
set_text(sh2[20],
    'No in-app chat, reactions, or community discussion. '
    'Event registration runs via Telegram bot + QR → Google Docs, entirely outside the app.'
)

# Current app — Static Profile description
set_text(sh2[23],
    'Basic info only; no certificate download, attendance history, achievements, skills, CV, '
    'or applied-jobs/grants tracking.'
)

# Prototype — Personalised Home Feed description (add Jobs + match %)
set_text(sh2[29],
    '5 category tabs (Events / Courses / Grants / Jobs / Podcasts), featured cards with coin-reward badges, '
    'match-% personalised ranking, weekly leaderboard rank on home header.'
)

# Prototype — Achievement Profile description (add Skills & CV, Jobs applied)
set_text(sh2[44],
    'Stats (coins / events attended / certs earned), Skills & CV section (skill tags + CV upload + AI gap analysis), '
    'Jobs & Internships applied, certificate PDF download, leaderboard rank promo.'
)

print('Slide 2 updated.')

# ─────────────────────────────────────────────
# SLIDE 3  —  Feature Gap Matrix
# ─────────────────────────────────────────────
s3 = prs.slides[2]
sh3 = s3.shapes

# Home & Discovery — prototype column (add Jobs tab + match %)
set_text(sh3[15],
    'Category tabs (Events / Courses / Grants / Jobs), featured cards, '
    'match-% personalised ranking, coin-reward labels, weekly rank on header'
)

# Gamification — current column (be specific about invisibility)
set_text(sh3[20],
    'ZiyoCoins exist in backend but are invisible in UI; '
    'no balance shown, no earning events, no redemption; no leaderboard anywhere in app'
)

# Content Formats — current column (add image distortion)
set_text(sh3[26],
    'Text/image cards only; no short-form video, no podcast player; '
    'Android image distortion bug ("smashed" images) reported for months, unfixed'
)

# Profile & History — prototype column (add Skills & CV)
set_text(sh3[39],
    'Stats card, Skills & CV (skill tags + CV upload + AI recommendations), '
    'certificate download, attended/upcoming tabs, Jobs applied, leaderboard promo'
)

print('Slide 3 updated.')

# ─────────────────────────────────────────────
# SLIDE 4  —  UX Assessment
# ─────────────────────────────────────────────
s4 = prs.slides[3]
sh4 = s4.shapes

# No engagement loop — description (add 83.8% metric)
set_text(sh4[37],
    '83.8 % Android uninstall rate (2,018 of 2,408 installs); iOS active users fell from 121 peak to 14; '
    'all growth is promotion-driven with no organic return flywheel'
)

# Fragmented content — description (add Aksana quote)
set_text(sh4[41],
    '"The app is pretty absent — there\'s nothing in it. Quite grey, not bright, doesn\'t attract users\' '
    'attention." — Aksana (UX Designer). Interface unchanged for 12+ months.'
)

# Absent personalisation — description (add designer context)
set_text(sh4[45],
    'Same grey UI for all users; no interest model; no onboarding. '
    'UX Designer cannot iterate — needs file access from developer (Azam) to make any changes.'
)

# Shallow social proof — description (add Telegram finding)
set_text(sh4[49],
    'Event registration via Telegram bot + QR → Google Docs, not in-app. '
    'No attendee counts, reactions, or post-event feedback loop visible to users.'
)

# Missing achievement system — description (add app store reviews)
set_text(sh4[53],
    'No certificates in profile; no Skills & CV; no Jobs/Grants applied tracking. '
    '5 unanswered 1-star app-store reviews — critical signal left without response.'
)

print('Slide 4 updated.')

# ─────────────────────────────────────────────
# SLIDE 5  —  Priority Recommendations
# ─────────────────────────────────────────────
s5 = prs.slides[4]
sh5 = s5.shapes

# Phase 1 — replace "Notification Centre" block with "Fix Android image distortion"
set_text(sh5[13], 'Fix Android image distortion')
set_text(sh5[15],
    "Images appear 'smashed' on Android — reported for months, still unfixed. "
    "Affects all 297 active Android users; critical for retention."
)

# Phase 1 — replace "Attendee count & social proof" block with "Reply to 1-star reviews"
set_text(sh5[17], 'Reply to 5 unanswered 1-star app-store reviews')
set_text(sh5[19],
    'Public response with V2 timeline commitment improves store ranking, '
    'rebuild trust, and signals active development.'
)

# Phase 1 — replace "Certificate download" block with "In-app event registration"
set_text(sh5[21], 'In-app event registration')
set_text(sh5[23],
    'Replace Telegram bot + QR → Google Docs flow with native registration. '
    'Top request from Events Head (Sirojiddin) — enables proper attendee tracking.'
)

print('Slide 5 updated.')

# ─────────────────────────────────────────────
# Save
# ─────────────────────────────────────────────
prs.save(PPTX_PATH)
print(f'\nSaved → {PPTX_PATH}')
